import os
import re
import shutil
from PIL import Image
import pytesseract
from docx import Document
from docx.shared import Pt, RGBColor
from docx.enum.text import WD_LINE_SPACING

def configurar_tesseract():
    rutas_posibles = [
        shutil.which('tesseract'),
        r'C:\Program Files\Tesseract-OCR\tesseract.exe',
        r'C:\Program Files (x86)\Tesseract-OCR\tesseract.exe',
    ]

    for ruta in rutas_posibles:
        if ruta and os.path.isfile(ruta):
            pytesseract.pytesseract.tesseract_cmd = ruta
            return

    raise RuntimeError(
        'Tesseract OCR no está instalado. Instálalo desde '
        'https://github.com/UB-Mannheim/tesseract/wiki y vuelve a ejecutar el script.'
    )


configurar_tesseract()

# Ruta exacta de tu carpeta de imágenes
CARPETA_IMAGENES = r"E:\AUTOMATIZACION_DE ARCHIVOS_2026\IMAGENES"
doc = Document()

# Configuración inicial del documento según tus especificaciones
style = doc.styles['Normal']
font = style.font
font.name = 'Times New Roman'
font.size = Pt(12)

# Obtener archivos de imagen y ordenarlos de forma natural (img 1, img 2... img 10)
def natural_sort_key(s):
    return [int(text) if text.isdigit() else text.lower() for text in re.split(r'(\d+)', s)]

try:
    archivos = [f for f in os.listdir(CARPETA_IMAGENES) if f.lower().endswith(('.png', '.jpg', '.jpeg', '.bmp', '.tiff'))]
    archivos.sort(key=natural_sort_key)
    
    print(f"Se seleccionaron {len(archivos)} imágenes. Procesando...")

    for idx, archivo in enumerate(archivos, start=1):
        ruta_imagen = os.path.join(CARPETA_IMAGENES, archivo)
        
        # Agregar encabezado numerado para cada imagen
        p_titulo = doc.add_paragraph()
        run_titulo = p_titulo.add_run(f"img {idx}")
        run_titulo.bold = True
        run_titulo.font.name = 'Times New Roman'
        run_titulo.font.size = Pt(12)
        
        try:
            # Extraer texto en inglés usando Tesseract OCR
            imagen = Image.open(ruta_imagen)
            # lang='eng' asegura la lectura correcta en inglés
            texto_extraido = pytesseract.image_to_string(imagen, lang='eng')
            
            # Conservar la salida completa del OCR, incluidos guiones y saltos de línea.
            if texto_extraido.strip():
                p = doc.add_paragraph()
                p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE

                run = p.add_run(texto_extraido.rstrip('\n'))
                run.font.name = 'Times New Roman'
                run.font.size = Pt(12)
                run.font.color.rgb = RGBColor(0, 0, 0)
                    
        except Exception as e:
            print(f"Error procesando {archivo}: {e}")
            
        # Añadir un salto de párrafo pequeño entre imágenes
        doc.add_paragraph()

    doc.save("resultado_198_imagenes.docx")
    print("¡Procesamiento completado con éxito! Archivo guardado como 'resultado_198_imagenes.docx'.")

except FileNotFoundError:
    print(f"No se encontró la ruta especificada: {CARPETA_IMAGENES}. Verifica que la carpeta exista.")
except Exception as e:
    print(f"Ocurrió un error: {e}")